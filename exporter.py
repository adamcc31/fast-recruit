"""
STEP 6 — Excel Exporter
Outputs three sheets:
  1. Ranked Candidates  — scored CVs, sorted by overall score, color-coded
  2. Failed Validation  — CVs that didn't pass the first gate
  3. Summary            — pipeline statistics
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config
from scorer import ScoreResult
from validator import ValidationResult


# ── Color map ──────────────────────────────────────────────────────────────────

FILL = {
    "SHORTLIST": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "REVIEW":    PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "REJECT":    PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "ERROR":     PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "HEADER":    PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid"),
}

THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 55)


def _style_header_row(ws):
    for cell in ws[1]:
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = FILL["HEADER"]
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border    = THIN_BORDER


def _find_col(ws, header_name: str) -> int | None:
    """Return 1-based column index of a header, or None."""
    for cell in ws[1]:
        if cell.value == header_name:
            return cell.column
    return None


def _knockout_label(r) -> str:
    """LOLOS, or GUGUR reason for knockout columns in Excel.

    Only reports criteria the user actually enabled in .env — stray model
    flags never produce a GUGUR label on their own.
    """
    from config import Config
    fails = []
    if Config.JOB_LOCATION and not getattr(r, "meets_location", True):
        fails.append(f"lokasi ({getattr(r, 'location', '') or 'Unknown'})")
    if Config.MIN_YEARS_EXPERIENCE > 0 and not getattr(r, "meets_experience", True):
        fails.append(f"pengalaman ({r.years_of_experience} th)")
    return "LOLOS" if not fails else "GUGUR: " + "; ".join(fails)


# ── Sheet builders ─────────────────────────────────────────────────────────────

def _build_ranked_sheet(score_results: list[ScoreResult], dimensions: list[str]) -> pd.DataFrame:
    rows = []
    for r in score_results:
        row = {
            "Nama File":         r.filename,
            "Nama Kandidat":     r.candidate_name,
            "Skor":              r.overall_score,
            "Rekomendasi":       r.recommendation,
            "Lokasi":            getattr(r, "location", "") or "—",
            "No. HP/WA":         getattr(r, "contact_phone", "") or "—",
            "Email":             getattr(r, "contact_email", "") or "—",
            "Knockout":          _knockout_label(r),
            "Pengalaman (th)":   r.years_of_experience,
            "Token Masuk":       getattr(r, "prompt_tokens", 0),
            "Token Keluar":      getattr(r, "completion_tokens", 0),
            "Token Reasoning":   getattr(r, "reasoning_tokens", 0),
            "Token Total":       getattr(r, "total_tokens", 0),
            "Skill Utama":       " | ".join(r.key_skills),
            "Ringkasan":         r.summary,
            "Kelebihan":         " | ".join(r.strengths),
            "Kekurangan":        " | ".join(r.red_flags) if r.red_flags else "—",
        }
        for dim in dimensions:
            dim_data = r.dimensions.get(dim, {})
            row[f"Skor {dim}"]  = dim_data.get("score", "—")
            row[f"Alasan {dim}"] = dim_data.get("reasoning", "—")
        rows.append(row)

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values("Skor", ascending=False).reset_index(drop=True)
    return df


def _build_failed_sheet(failed: list[ValidationResult]) -> pd.DataFrame:
    rows = []
    for v in failed:
        rows.append({
            "Nama File":       v.filename,
            "Jenis Gagal":     v.status,
            "Alasan":          v.reason,
            "Skor Struktur":   f"{v.structure_score:.0%}",
            "Ada Email":       "✓" if v.has_email else "✗",
            "Ada No. HP":      "✓" if v.has_phone else "✗",
            "Ada Tanggal":     "✓" if v.has_dates else "✗",
            "Ada Pendidikan":  "✓" if v.has_education else "✗",
            "Ada Pengalaman":  "✓" if v.has_experience else "✗",
        })
    return pd.DataFrame(rows)


def _build_summary_sheet(
    total: int,
    scored: list[ScoreResult],
    failed: list[ValidationResult],
) -> pd.DataFrame:
    shortlisted = sum(1 for r in scored if r.recommendation == "SHORTLIST")
    reviewed    = sum(1 for r in scored if r.recommendation == "REVIEW")
    rejected    = sum(1 for r in scored if r.recommendation == "REJECT")
    errors      = sum(1 for r in scored if r.recommendation == "ERROR")

    fail_extraction = sum(1 for v in failed if v.status == "FAIL_EXTRACTION")
    fail_structure  = sum(1 for v in failed if v.status == "FAIL_STRUCTURE")
    fail_content    = sum(1 for v in failed if v.status == "FAIL_CONTENT")
    fail_non_cv     = sum(1 for v in failed if v.status == "FAIL_NON_CV")

    scored_only = [r.overall_score for r in scored if r.recommendation != "ERROR"]
    avg_score = round(sum(scored_only) / len(scored_only), 2) if scored_only else 0

    tok_in  = sum(getattr(r, "prompt_tokens", 0) for r in scored)
    tok_out = sum(getattr(r, "completion_tokens", 0) for r in scored)
    tok_rs  = sum(getattr(r, "reasoning_tokens", 0) for r in scored)
    tok_all = sum(getattr(r, "total_tokens", 0) for r in scored)

    rows = [
        {"Metric": "── Input ───────────────", "Value": ""},
        {"Metric": "Total CV",                "Value": total},
        {"Metric": "Lolos Validasi",          "Value": len(scored)},
        {"Metric": "Gagal Validasi",          "Value": len(failed)},
        {"Metric": "── Rincian Gagal ───────", "Value": ""},
        {"Metric": "  FAIL_EXTRACTION",       "Value": fail_extraction},
        {"Metric": "  FAIL_STRUCTURE",        "Value": fail_structure},
        {"Metric": "  FAIL_CONTENT",          "Value": fail_content},
        {"Metric": "  FAIL_NON_CV",           "Value": fail_non_cv},
        {"Metric": "── Hasil Scoring ───────", "Value": ""},
        {"Metric": "SHORTLIST",               "Value": shortlisted},
        {"Metric": "REVIEW",                  "Value": reviewed},
        {"Metric": "REJECT",                  "Value": rejected},
        {"Metric": "ERROR (gagal LLM/parse)", "Value": errors},
        {"Metric": "Rata-rata Skor",          "Value": avg_score},
        {"Metric": "── Token ───────────────", "Value": ""},
        {"Metric": "Token Masuk (prompt)",    "Value": tok_in},
        {"Metric": "Token Keluar",            "Value": tok_out},
        {"Metric": "Token Reasoning",         "Value": tok_rs},
        {"Metric": "Token Total",             "Value": tok_all},
    ]

    if Config.LLM_COST_INPUT_PER_1M > 0 or Config.LLM_COST_OUTPUT_PER_1M > 0:
        est = tok_in / 1_000_000 * Config.LLM_COST_INPUT_PER_1M \
            + tok_out / 1_000_000 * Config.LLM_COST_OUTPUT_PER_1M
        rows.append({"Metric": "Estimasi Biaya (USD)", "Value": round(est, 4)})

    return pd.DataFrame(rows)


# ── Main export function ───────────────────────────────────────────────────────

def export_to_excel(
    score_results: list[ScoreResult],
    failed_validations: list[ValidationResult],
    output_path: Path,
    dimensions: list[str],
) -> None:
    """Write all results to a formatted multi-sheet Excel file."""

    total = len(score_results) + len(failed_validations)

    df_ranked  = _build_ranked_sheet(score_results, dimensions)
    df_failed  = _build_failed_sheet(failed_validations)
    df_summary = _build_summary_sheet(total, score_results, failed_validations)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(str(output_path), engine="openpyxl") as writer:
        df_ranked.to_excel( writer, sheet_name="Ranking Kandidat", index=False)
        df_failed.to_excel( writer, sheet_name="Gagal Validasi",   index=False)
        df_summary.to_excel(writer, sheet_name="Ringkasan",        index=False)

        wb = writer.book

        # ── Format: Ranking Kandidat ───────────────────────────
        ws = wb["Ranking Kandidat"]
        _style_header_row(ws)
        _auto_width(ws)

        rec_col = _find_col(ws, "Rekomendasi")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            rec_val = row[rec_col - 1].value if rec_col else None
            fill    = FILL.get(rec_val, FILL["ERROR"])
            for cell in row:
                cell.fill      = fill
                cell.border    = THIN_BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # ── Format: Gagal Validasi ─────────────────────────────
        ws2 = wb["Gagal Validasi"]
        _style_header_row(ws2)
        _auto_width(ws2)
        for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
            for cell in row:
                cell.border    = THIN_BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # ── Format: Ringkasan ──────────────────────────────────
        ws3 = wb["Ringkasan"]
        _style_header_row(ws3)
        _auto_width(ws3)
        for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
            for cell in row:
                cell.border = THIN_BORDER
